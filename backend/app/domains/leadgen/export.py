"""Выгрузка базы компаний в Excel: одна строка = одна компания, все контакты и контекст рекламы."""
from __future__ import annotations

import io
from collections import defaultdict

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.domains.leadgen.models import LgAd, LgCompany, LgContact

MAX_ROWS = 5000
STAGE_NAMES = {"new": "Новая", "enriched": "Обогащена", "ready": "Готова", "in_outreach": "В аутриче",
               "replied": "Ответила", "converted": "Клиент", "rejected": "Отказ", "dnc": "Стоп-лист"}
FIT_NAMES = {"fit": "Подходит", "maybe": "Возможно", "no": "Не подходит"}
LEGAL_NAMES = {"active": "Действует", "liquidated": "Ликвидирована / банкротство"}

COLUMNS = [
    ("Компания", 28), ("Сайт", 26), ("Город", 14), ("Ниша", 18), ("Скоринг", 9), ("Почему такой скоринг", 40),
    ("Оценка ЦА", 12), ("Причина (ИИ)", 34), ("Стадия", 12),
    ("Телефоны", 22), ("WhatsApp", 18), ("Telegram", 18), ("Email", 28), ("Личные email (не писать)", 26),
    ("ИНН", 14), ("Юрлицо", 30), ("Руководитель", 26), ("Статус юрлица", 16), ("Выручка, ₽", 14),
    ("Что рекламирует", 44), ("Текст объявления", 60), ("Ключей с показами", 10), ("Спецразмещение", 10),
    ("Впервые замечена", 12), ("Карточка в портале", 30),
]


def _join(values) -> str:
    return ", ".join(dict.fromkeys(v for v in values if v))


async def build_xlsx(db: AsyncSession, companies: list[LgCompany], *, base_url: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ids = [c.id for c in companies]
    contacts: dict[int, list[LgContact]] = defaultdict(list)
    ads: dict[int, list[LgAd]] = defaultdict(list)
    for i in range(0, len(ids), 500):
        chunk = ids[i:i + 500]
        for row in (await db.execute(select(LgContact).where(LgContact.company_id.in_(chunk))
                                     .order_by(LgContact.is_primary.desc(), LgContact.id))).scalars():
            contacts[row.company_id].append(row)
        for row in (await db.execute(select(LgAd).where(LgAd.company_id.in_(chunk))
                                     .order_by(LgAd.last_seen_at.desc()))).scalars():
            ads[row.company_id].append(row)

    wb = Workbook()
    ws = wb.active
    ws.title = "Компании"
    ws.append([name for name, _ in COLUMNS])
    head_fill = PatternFill("solid", fgColor="1F3A68")
    for idx, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 32

    portal = (base_url or "").rstrip("/")
    for c in companies:
        cs = [x for x in contacts[c.id] if not x.bounced]
        work_email = [x.value_norm for x in cs if x.kind == "email" and not x.is_personal]
        personal = [x.value_norm for x in cs if x.kind == "email" and x.is_personal]
        ad_list = ads[c.id]
        top = ad_list[0] if ad_list else None
        keywords = {k for a in ad_list for k in (a.keywords or [])}
        ws.append([
            c.display_name or c.domain,
            c.domain,
            c.city,
            c.niche,
            c.score,
            "; ".join(f"{r.get('label')} ({'+' if r.get('weight', 0) > 0 else ''}{r.get('weight')})"
                      for r in (c.score_reasons or [])),
            FIT_NAMES.get(c.fit_label or "", ""),
            c.fit_reason,
            STAGE_NAMES.get(c.stage, c.stage),
            _join(x.value_norm for x in cs if x.kind == "phone"),
            _join(x.value_norm for x in cs if x.kind == "whatsapp"),
            _join(("@" + x.value_norm.lstrip("@")) for x in cs if x.kind == "telegram"),
            _join(work_email),
            _join(personal),
            c.inn,
            c.legal_name,
            " — ".join(v for v in (c.director_name, c.director_post) if v) or None,
            LEGAL_NAMES.get(c.legal_status or "", c.legal_status),
            c.revenue_rub,
            top.title if top else None,
            top.text if top else None,
            len(keywords) or None,
            "да" if any(a.placement == "premium" for a in ad_list) else None,
            c.first_seen_at.date() if c.first_seen_at else None,
            f"{portal}/admin/leadgen/companies?company={c.id}" if portal else None,
        ])
        r = ws.max_row
        if c.domain:
            site = ws.cell(row=r, column=2)
            site.hyperlink, site.style = f"https://{c.domain}", "Hyperlink"
        if portal:
            card = ws.cell(row=r, column=len(COLUMNS))
            card.hyperlink, card.style = card.value, "Hyperlink"
        for col in (6, 8, 20, 21):
            ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=24).number_format = "DD.MM.YYYY"
        ws.cell(row=r, column=19).number_format = "# ##0"

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
