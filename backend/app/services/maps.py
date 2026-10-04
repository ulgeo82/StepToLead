"""Yandex Maps (Яндекс Бизнес) and 2GIS as advertising channels.

Neither service has a public statistics API (the Yandex Business API only creates and pays subscriptions), so a map
channel is fed three ways and then counts in analytics like Direct / VK / Avito:
  * spend — the monthly promotion subscription, spread evenly over the days of the month;
  * card statistics — the Excel/CSV export from the cabinet (or monthly totals typed in): views, clicks into the card,
    "call" taps, routes, site clicks (stored on AdMetricDaily: impressions, clicks, raw{calls, routes, site});
  * leads — calls to a call-tracking number placed in the card (telephony line_map → this connection) and site visits
    with the channel's utm_source (matched server-side in inbound_lead).
"""
import calendar
import csv
import io
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.marketing import AdConnection, AdMetricDaily

PLATFORMS = {"yandex_maps": "Яндекс Карты", "2gis": "2ГИС"}
DEFAULT_UTM = {"yandex_maps": "yandex_maps", "2gis": "2gis"}
EXTRA = ("calls", "routes", "site")
MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "мая": 5, "май": 5, "июн": 6, "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12}


class StatsError(ValueError):
    pass


def is_map(row: AdConnection | None) -> bool:
    return bool(row and row.platform in PLATFORMS)


def utm_source(row: AdConnection) -> str:
    return str((row.config or {}).get("utm_source") or DEFAULT_UTM.get(row.platform, row.platform)).lower()


# --------------------------------------------------------------------------- statistics file

def classify(header: str) -> str | None:
    h = header.lower().replace("ё", "е").strip()
    if not h:
        return None
    if h in {"дата", "день", "период", "date"} or h.startswith("дата"):
        return "date"
    if "расход" in h or "списан" in h or "стоимост" in h or "потрачен" in h:
        return "spend"
    if "маршрут" in h:
        return "routes"
    if "позвон" in h or "звонк" in h or "телефон" in h:
        return "calls"
    if "сайт" in h:
        return "site"
    if "переход" in h or "открыти" in h or "клик" in h or "просмотр карточ" in h or "в профиль" in h:
        return "clicks"
    if "показ" in h or "просмотр" in h:
        return "impressions"
    return None


def parse_day(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip().lower()
    if not text:
        return None
    text = re.split(r"\s+[–—-]\s+|\s*[–—]\s*", text)[0].strip()  # "01.10.2026 – 07.10.2026" → start of the range
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    match = re.match(r"(\d{1,2})\s+([а-я]+)\.?\s+(\d{4})", text)
    if match:
        month = next((n for key, n in MONTHS.items() if match.group(2).startswith(key)), None)
        if month:
            try:
                return date(int(match.group(3)), month, int(match.group(1)))
            except ValueError:
                return None
    return None


def number(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", "").replace(" ", "").replace("₽", "").replace(",", ".")
    try:
        return float(text) if text not in {"", "-", "—"} else None
    except ValueError:
        return None


def tables(filename: str, content: bytes) -> list[list[list]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [[list(r) for r in sheet.iter_rows(values_only=True)] for sheet in book.worksheets]
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise StatsError("Не удалось прочитать кодировку CSV")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[list(r) for r in csv.reader(io.StringIO(text), dialect)]]
    raise StatsError("Загрузите выгрузку статистики в формате XLSX или CSV")


def parse_stats(filename: str, content: bytes) -> tuple[dict[date, dict], list[str]]:
    """Every sheet: find a header row with a date column and known metrics; sum the values per day."""
    days: dict[date, dict] = {}
    found: set[str] = set()
    for rows in tables(filename, content):
        header_index, mapping = None, {}
        for i, row in enumerate(rows[:15]):
            kinds = {j: classify(str(cell or "")) for j, cell in enumerate(row)}
            if "date" in kinds.values() and len({k for k in kinds.values() if k and k != "date"}) >= 1:
                header_index, mapping = i, {j: k for j, k in kinds.items() if k}
                break
        if header_index is None:
            continue
        date_col = next(j for j, k in mapping.items() if k == "date")
        for row in rows[header_index + 1:]:
            day = parse_day(row[date_col] if date_col < len(row) else None)
            if not day:
                continue
            bucket = days.setdefault(day, {})
            for j, kind in mapping.items():
                if kind == "date" or j >= len(row):
                    continue
                value = number(row[j])
                if value is not None:
                    bucket[kind] = bucket.get(kind, 0) + value
                    found.add(kind)
    if not days:
        raise StatsError("В файле не нашли таблицу с датами. Выгрузите статистику по дням: Яндекс Бизнес → Статистика → значок загрузки у графика или «Собрать статистику»")
    return days, sorted(found)


# --------------------------------------------------------------------------- storage

async def metric_rows(db: AsyncSession, conn: AdConnection, start: date, end: date) -> dict[date, AdMetricDaily]:
    rows = (await db.scalars(select(AdMetricDaily).where(AdMetricDaily.connection_id == conn.id, AdMetricDaily.date >= start,
                                                         AdMetricDaily.date <= end))).all()
    return {r.date: r for r in rows}


def row_for(db: AsyncSession, existing: dict, conn: AdConnection, day: date) -> AdMetricDaily:
    row = existing.get(day)
    if row is None:
        row = AdMetricDaily(connection_id=conn.id, date=day, spend=0, impressions=0, clicks=0, leads=0, raw={})
        db.add(row)
        existing[day] = row
    return row


def month_bounds(month: str) -> tuple[date, date]:
    year, mon = (int(x) for x in month.split("-"))
    return date(year, mon, 1), date(year, mon, calendar.monthrange(year, mon)[1])


def spread(total: float, days: int, cents: bool) -> list[float]:
    """Even split whose parts add up exactly to the total (money in kopecks, counts in whole units)."""
    unit = 100 if cents else 1
    whole = round(total * unit)
    base, rest = divmod(whole, days)
    return [(base + (1 if i < rest else 0)) / unit for i in range(days)]


async def set_month(db: AsyncSession, conn: AdConnection, month: str, *, budget: float | None = None,
                    totals: dict | None = None) -> None:
    """Monthly subscription cost and/or monthly totals typed in by hand. Does not commit."""
    start, end = month_bounds(month)
    existing = await metric_rows(db, conn, start, end)
    n = (end - start).days + 1
    config = dict(conn.config or {})
    if budget is not None:
        budgets = dict(config.get("budgets") or {})
        if budget > 0:
            budgets[month] = budget
        else:
            budgets.pop(month, None)
        config["budgets"] = budgets
        for i, part in enumerate(spread(budget, n, True)):
            row = row_for(db, existing, conn, date.fromordinal(start.toordinal() + i))
            row.spend = part
    if totals:
        manual = dict(config.get("manual") or {})
        manual[month] = {k: int(v) for k, v in totals.items() if v is not None}
        config["manual"] = manual
        for key, value in totals.items():
            if value is None:
                continue
            for i, part in enumerate(spread(float(value), n, False)):
                row = row_for(db, existing, conn, date.fromordinal(start.toordinal() + i))
                if key in {"impressions", "clicks"}:
                    setattr(row, key, int(part))
                else:
                    row.raw = {**(row.raw or {}), key: int(part)}
    conn.config = config


async def import_days(db: AsyncSession, conn: AdConnection, days: dict[date, dict]) -> dict:
    """Store parsed daily statistics. Spend from the file is used only for months without a typed-in budget. Does not commit."""
    if not days:
        return {"days": 0}
    budgets = (conn.config or {}).get("budgets") or {}
    existing = await metric_rows(db, conn, min(days), max(days))
    for day, values in days.items():
        row = row_for(db, existing, conn, day)
        if "impressions" in values:
            row.impressions = int(values["impressions"])
        if "clicks" in values:
            row.clicks = int(values["clicks"])
        if "spend" in values and day.strftime("%Y-%m") not in budgets:
            row.spend = round(values["spend"], 2)
        extra = {k: int(values[k]) for k in EXTRA if k in values}
        if extra:
            row.raw = {**(row.raw or {}), **extra}
    return {"days": len(days), "first_date": min(days).isoformat(), "last_date": max(days).isoformat()}


async def totals(db: AsyncSession, conn: AdConnection, start: date, end: date) -> dict:
    rows = (await metric_rows(db, conn, start, end)).values()
    out = {"spend": 0.0, "impressions": 0, "clicks": 0, **{k: 0 for k in EXTRA}}
    for r in rows:
        out["spend"] += float(r.spend or 0)
        out["impressions"] += int(r.impressions or 0)
        out["clicks"] += int(r.clicks or 0)
        for k in EXTRA:
            out[k] += int((r.raw or {}).get(k) or 0)
    out["spend"] = round(out["spend"], 2)
    out["has_data"] = any(int(r.impressions or 0) or int(r.clicks or 0) or any((r.raw or {}).get(k) for k in EXTRA) for r in rows)
    return out


async def connection_for_utm(db: AsyncSession, project_id: int, source: str | None) -> int | None:
    """Server-side: a lead whose utm_source is a map channel's marker belongs to that channel."""
    value = (source or "").strip().lower()
    if not value:
        return None
    rows = (await db.scalars(select(AdConnection).where(AdConnection.project_id == project_id,
                                                        AdConnection.platform.in_(list(PLATFORMS)),
                                                        AdConnection.status != "disconnected"))).all()
    return next((r.id for r in rows if utm_source(r) == value), None)
