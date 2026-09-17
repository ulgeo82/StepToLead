import csv
import io
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class ImportedLead:
    username: str
    first_message: str


def _normalize_headers(row: list[object]) -> list[str]:
    return [str(value or "").strip().lower() for value in row]


def _to_leads(rows: list[list[object]]) -> tuple[list[ImportedLead], list[str]]:
    if not rows:
        raise ValueError("Файл пуст")
    headers = _normalize_headers(rows[0])
    required = {"username", "first_message"}
    if not required.issubset(headers):
        raise ValueError("Нужны колонки username и first_message")

    username_index = headers.index("username")
    message_index = headers.index("first_message")
    leads: list[ImportedLead] = []
    errors: list[str] = []
    seen: set[str] = set()

    for line_number, row in enumerate(rows[1:], start=2):
        values = list(row) + [""] * (len(headers) - len(row))
        username = str(values[username_index] or "").strip().lstrip("@").lower()
        message = str(values[message_index] or "").strip()
        if not username or not message:
            errors.append(f"Строка {line_number}: пустой username или first_message")
            continue
        if username in seen:
            errors.append(f"Строка {line_number}: дубликат @{username}")
            continue
        seen.add(username)
        leads.append(ImportedLead(username=username, first_message=message))
    return leads, errors


def parse_leads_file(filename: str, content: bytes) -> tuple[list[ImportedLead], list[str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;")
        except csv.Error:
            dialect = csv.excel
        rows = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
        return _to_leads(rows)
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        return _to_leads(rows)
    raise ValueError("Поддерживаются только CSV и XLSX")

