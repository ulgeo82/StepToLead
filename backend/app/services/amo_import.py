"""Move from amoCRM: import the deals export (Сделки → «Экспорт» → Excel/CSV) into the StepToLead CRM.

Two steps: preview (columns found, amo stages and managers with suggested matches) and import with the chosen mapping.
Contacts are merged by phone/email, deals keep their amo creation date, won deals can become confirmed sales
(no offline conversions are sent for history). Re-importing the same file skips deals imported before (amo ID).
"""
import csv
import io
import re
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

MAX_ROWS = 5000
COLUMNS = {
    "amo_id": ("id", "ид", "id сделки"),
    "name": ("название сделки", "сделка", "название"),
    "amount": ("бюджет", "бюджет ₽", "бюджет, ₽", "сумма", "бюджет руб"),
    "stage": ("этап сделки", "этап", "статус", "статус сделки"),
    "pipeline": ("воронка", "название воронки"),
    "responsible": ("ответственный", "ответственный пользователь", "ответственный за сделку"),
    "created": ("дата создания", "создана", "дата создания сделки"),
    "closed": ("дата закрытия", "закрыта", "дата закрытия сделки"),
    "tags": ("теги", "теги сделки"),
    "contact": ("основной контакт", "контакт", "полное имя контакта", "имя контакта", "фио"),
    "company": ("компания", "название компании", "компания контакта"),
    "note": ("примечание", "примечания", "комментарий", "описание"),
    "utm_source": ("utm_source",), "utm_medium": ("utm_medium",), "utm_campaign": ("utm_campaign",),
    "utm_content": ("utm_content",), "utm_term": ("utm_term",),
}
WON_WORDS = ("успешно реализовано", "успешно", "выиграна", "продажа", "оплачено")
LOST_WORDS = ("закрыто и не реализовано", "не реализовано", "отказ", "проиграна", "потеряна")


class ImportError_(ValueError):
    pass


def read_rows(filename: str, content: bytes) -> list[list[object]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".xlsx":
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ImportError_("Не удалось прочитать кодировку CSV")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    raise ImportError_("Загрузите выгрузку amoCRM в формате XLSX или CSV")


def header_map(headers: list[str]) -> tuple[dict[str, int], list[int], list[int]]:
    found: dict[str, int] = {}
    low = [h.strip().lower().replace("ё", "е") for h in headers]
    for field, aliases in COLUMNS.items():
        for alias in aliases:
            if alias in low and field not in found:
                found[field] = low.index(alias)
    phones = [i for i, h in enumerate(low) if "телефон" in h]
    emails = [i for i, h in enumerate(low) if "email" in h or "e-mail" in h or "почта" in h]
    return found, phones, emails


def parse_date(value) -> datetime | None:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=timezone.utc)
    text = str(value or "").strip()
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def parse_amount(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value) if value > 0 else None
    text = re.sub(r"[^\d,.]", "", str(value or "")).replace(",", ".")
    try:
        number = float(text) if text else None
    except ValueError:
        return None
    return number if number and number > 0 else None


def split_values(value) -> list[str]:
    return [v.strip() for v in re.split(r"[,;\n]", str(value or "")) if v.strip()]


def parse(filename: str, content: bytes) -> tuple[list[dict], dict]:
    rows = [r for r in read_rows(filename, content) if any(str(c or "").strip() for c in r)]
    if len(rows) < 2:
        raise ImportError_("В файле нет сделок")
    headers = [str(h or "") for h in rows[0]]
    found, phone_cols, email_cols = header_map(headers)
    if "name" not in found and "contact" not in found:
        raise ImportError_("Не нашли колонки «Название сделки» или «Основной контакт» — это точно выгрузка сделок amoCRM?")
    if len(rows) - 1 > MAX_ROWS:
        raise ImportError_(f"В файле {len(rows) - 1} строк — за раз можно загрузить до {MAX_ROWS}. Разбейте выгрузку по периодам.")

    def cell(row, field):
        index = found.get(field)
        return row[index] if index is not None and index < len(row) else None

    deals = []
    for row in rows[1:]:
        phones, emails = [], []
        for i in phone_cols:
            phones += [p for p in split_values(row[i] if i < len(row) else "") if len(re.sub(r"\D", "", p)) >= 10]
        for i in email_cols:
            emails += [e.lower() for e in split_values(row[i] if i < len(row) else "") if "@" in e]
        contact = str(cell(row, "contact") or "").strip()
        name = str(cell(row, "name") or "").strip()
        deals.append({
            "amo_id": str(cell(row, "amo_id") or "").strip()[:40] or None,
            "name": (name or (f"Сделка: {contact}" if contact else "Сделка из amoCRM"))[:240],
            "amount": parse_amount(cell(row, "amount")),
            "stage": str(cell(row, "stage") or "").strip()[:120] or "Без этапа",
            "pipeline": str(cell(row, "pipeline") or "").strip()[:120] or None,
            "responsible": str(cell(row, "responsible") or "").strip()[:120] or None,
            "created": parse_date(cell(row, "created")), "closed": parse_date(cell(row, "closed")),
            "tags": [t[:40] for t in split_values(cell(row, "tags"))][:10],
            "contact": (contact or str(cell(row, "company") or "").strip() or name or "Клиент")[:180],
            "company": str(cell(row, "company") or "").strip()[:180] or None,
            "phones": list(dict.fromkeys(phones))[:3], "emails": list(dict.fromkeys(emails))[:3],
            "note": str(cell(row, "note") or "").strip()[:3000] or None,
            "utm": {k: str(cell(row, k)).strip()[:255] for k in ("utm_source", "utm_medium", "utm_campaign", "utm_content", "utm_term") if cell(row, k)},
        })
    columns = {field: headers[i] for field, i in found.items()}
    columns["phones"] = ", ".join(headers[i] for i in phone_cols) or None
    columns["emails"] = ", ".join(headers[i] for i in email_cols) or None
    return deals, columns


def suggest_stage(amo_stage: str, stages: list) -> int | None:
    """stages: CrmStage rows of the target pipeline (active)."""
    text = amo_stage.lower()
    by_type = lambda kind: next((s.id for s in stages if s.analytics_type == kind), None)
    if any(w in text for w in WON_WORDS):
        return by_type("WON")
    if any(w in text for w in LOST_WORDS):
        return by_type("LOST")
    exact = next((s.id for s in stages if s.name.strip().lower() == text), None)
    if exact:
        return exact
    loose = next((s.id for s in stages if s.analytics_type not in {"WON", "LOST"} and (s.name.lower() in text or text in s.name.lower())), None)
    return loose or by_type("LEAD")


def suggest_user(amo_name: str, users: dict[int, str]) -> int | None:
    """Exact name, then all words of one name inside the other («Анна» ↔ «Анна Ковалёва»); ambiguous → None."""
    norm = lambda v: " ".join((v or "").lower().replace("ё", "е").split())
    text = norm(amo_name)
    exact = [uid for uid, name in users.items() if norm(name) == text]
    if exact:
        return exact[0]
    words = set(text.split())
    loose = [uid for uid, name in users.items() if norm(name) and (set(norm(name).split()) <= words or words <= set(norm(name).split()))]
    return loose[0] if len(loose) == 1 else None


def summary(deals: list[dict], stages: list, users: dict[int, str]) -> dict:
    stage_counts = Counter(d["stage"] for d in deals)
    owner_counts = Counter(d["responsible"] for d in deals if d["responsible"])
    return {
        "total": len(deals), "no_contact": sum(1 for d in deals if not d["phones"] and not d["emails"]),
        "with_amount": sum(1 for d in deals if d["amount"]),
        "pipelines": sorted({d["pipeline"] for d in deals if d["pipeline"]}),
        "stages": [{"name": name, "count": count, "suggested": suggest_stage(name, stages)} for name, count in stage_counts.most_common()],
        "users": [{"name": name, "count": count, "suggested": suggest_user(name, users)} for name, count in owner_counts.most_common()],
        "sample": [{k: d[k] for k in ("name", "contact", "amount", "stage", "responsible")} | {"phone": (d["phones"] or [None])[0],
                    "created": d["created"].isoformat() if d["created"] else None} for d in deals[:5]],
    }
