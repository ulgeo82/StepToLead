"""Списки лидогенерации: список на каждый запуск, ручные списки, выгрузка по спискам «только новые», история."""
import io
import unittest

from openpyxl import load_workbook

import test_leadgen_api as base


class LeadgenListsTests(unittest.TestCase):
    setUp, tearDown = base.LeadgenApiTests.setUp, base.LeadgenApiTests.tearDown
    post, run_search = base.LeadgenApiTests.post, base.LeadgenApiTests.run_search

    def sheet(self, response):
        self.assertEqual(response.status_code, 200, response.text)
        ws = load_workbook(io.BytesIO(response.content)).active
        header = [c.value for c in ws[1]]
        return header, [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]

    def test_run_creates_list_and_export_only_new(self):
        run_id = self.run_search()
        lists = self.client.get("/api/admin/leadgen/lists").json()["items"]
        self.assertEqual(len(lists), 1)
        lst = lists[0]
        self.assertEqual((lst["run_id"], lst["kind"], lst["total"], lst["new"], lst["not_exported"]), (run_id, "run", 2, 2, 2))
        self.assertIn("Кухни на заказ · Екатеринбург", lst["name"])

        page = self.client.get("/api/admin/leadgen/companies", params={"list_id": lst["id"]}).json()
        self.assertEqual(page["total"], 2)
        self.assertEqual(page["items"][0]["lists"][0]["id"], lst["id"])

        count = self.client.get("/api/admin/leadgen/companies/export",
                                params={"list_id": lst["id"], "not_exported": True, "count_only": True}).json()
        self.assertEqual(count["count"], 2)
        header, rows = self.sheet(self.client.get("/api/admin/leadgen/companies/export",
                                                  params={"list_id": lst["id"], "not_exported": True, "preset": "outreach"}))
        self.assertEqual(len(rows), 2)
        self.assertNotIn("Скоринг", header); self.assertIn("Телефоны", header)
        # Второй раз «только новые» — уже пусто; история выгрузок запомнила файл и умеет его повторить.
        again = self.client.get("/api/admin/leadgen/companies/export",
                                params={"list_id": lst["id"], "not_exported": True, "count_only": True}).json()
        self.assertEqual(again["count"], 0)
        history = self.client.get("/api/admin/leadgen/exports").json()["items"]
        self.assertEqual((len(history), history[0]["count"], history[0]["preset"]), (1, 2, "outreach"))
        self.assertTrue(history[0]["title"].startswith("Список: "))
        _, repeated = self.sheet(self.client.get(f"/api/admin/leadgen/exports/{history[0]['id']}/download"))
        self.assertEqual(len(repeated), 2)
        self.assertEqual(len(self.client.get("/api/admin/leadgen/exports").json()["items"]), 1)  # повтор не пишет историю
        self.assertEqual(self.client.get("/api/admin/leadgen/lists").json()["items"][0]["not_exported"], 0)

    def test_manual_list_merge_archive_delete(self):
        self.run_search()
        ids = [c["id"] for c in self.client.get("/api/admin/leadgen/companies").json()["items"]]
        manual = self.post("/api/admin/leadgen/lists", {"name": "Горячие", "company_ids": ids[:1]}).json()
        self.assertEqual(manual["total"], 1)
        added = self.post(f"/api/admin/leadgen/lists/{manual['id']}/companies", {"company_ids": ids}).json()
        self.assertEqual((added["added"], added["total"]), (1, 2))
        removed = self.post(f"/api/admin/leadgen/lists/{manual['id']}/companies/remove", {"company_ids": ids[1:]}).json()
        self.assertEqual(removed["total"], 1)

        run_list = [l for l in self.client.get("/api/admin/leadgen/lists").json()["items"] if l["kind"] == "run"][0]
        merged = self.post("/api/admin/leadgen/lists/merge", {"list_ids": [run_list["id"], manual["id"]],
                                                               "name": "Все кухни"}).json()
        self.assertEqual(merged["total"], 2)  # дубли склеились
        active = self.client.get("/api/admin/leadgen/lists").json()
        self.assertEqual(([l["name"] for l in active["items"]], active["archived_count"]), (["Все кухни"], 2))
        archived = self.client.get("/api/admin/leadgen/lists", params={"archived": True}).json()["items"]
        self.assertEqual(len(archived), 2)
        restored = self.client.patch(f"/api/admin/leadgen/lists/{manual['id']}", json={"archived": False, "name": "Горячие 2"},
                                     headers=base.ORIGIN).json()
        self.assertEqual((restored["archived"], restored["name"]), (False, "Горячие 2"))

        # Удаление списка с компаниями: удаляются только компании, которых нет в других списках.
        solo = self.post("/api/admin/leadgen/lists", {"name": "Пробный"}).json()
        r = self.post(f"/api/admin/leadgen/lists/{solo['id']}/delete", {"with_companies": True}).json()
        self.assertEqual(r["deleted"], 0)
        r = self.post(f"/api/admin/leadgen/lists/{merged['id']}/delete", {"with_companies": True}).json()
        self.assertEqual((r["deleted"], r["kept_in_other_lists"]), (0, 2))
        self.assertEqual(self.client.get("/api/admin/leadgen/companies").json()["total"], 2)

    def test_past_runs_get_lists_and_merge_keeps_membership(self):
        import asyncio
        from sqlalchemy import delete
        from app.domains.leadgen.models import LgList, LgListItem
        self.run_search()

        async def wipe():  # как на проде до обновления: запуски есть, списков нет
            async with self.sessions() as db:
                await db.execute(delete(LgListItem)); await db.execute(delete(LgList)); await db.commit()
        asyncio.run(wipe())
        lists = self.client.get("/api/admin/leadgen/lists").json()["items"]
        self.assertEqual((len(lists), lists[0]["total"]), (1, 2))
        a, b = [c["id"] for c in self.client.get("/api/admin/leadgen/companies").json()["items"]]
        self.assertEqual(self.post("/api/admin/leadgen/companies/merge", {"keep_id": a, "drop_id": b}).status_code, 200)
        self.assertEqual(self.client.get("/api/admin/leadgen/lists").json()["items"][0]["total"], 1)


if __name__ == "__main__":
    unittest.main()
