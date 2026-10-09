"""API лидогенерации: доступ, запуск поиска, база компаний, сегменты, стоп-лист, настройки."""
import asyncio
import os
import tempfile
import unittest

import httpx
from datetime import datetime, timezone

from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api.router import api_router
from app.core.access import COOKIE, hash_password, token_digest
from app.core.config import settings
from app.db import Base, get_db
from app.domains.leadgen import routes
from app.domains.leadgen.core.serp import SerpAd
from app.domains.leadgen.providers import ProviderNotConfigured, StaticProvider
from app.models.access import AdminSession, AdminUser
from app.models.marketing import ClientWorkspace

ORIGIN = {"Origin": "http://localhost:3000"}
KEYS = ["кухни на заказ", "кухни под заказ", "заказать кухню"]
SEVER = SerpAd(keyword="", url="https://kuhni-sever.ru/", title="Кухни от 14 дней", text="Замер", premium=True)
DUB = SerpAd(keyword="", url="https://dub-mebel.ru/", title="Кухня мечты", text="Гарантия")
DATA = {KEYS[0]: [SEVER, DUB], KEYS[1]: [SEVER], KEYS[2]: [SEVER]}


class LeadgenApiTests(unittest.TestCase):
    def setUp(self):
        fd, self.path = tempfile.mkstemp(suffix=".sqlite"); os.close(fd)
        self.engine = create_async_engine(f"sqlite+aiosqlite:///{self.path}")
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)

        async def seed():
            async with self.engine.begin() as conn:
                await conn.run_sync(Base.metadata.create_all)
            async with self.sessions() as db:
                db.add(ClientWorkspace(id=7, name="StepToLead Agency"))
                db.add(AdminUser(id=1, username="admin", password_hash=hash_password("x" * 16), role="admin"))
                db.add(AdminSession(token_hash=token_digest("a" * 43), user_id=1,
                                    expires_at=datetime(2099, 1, 1, tzinfo=timezone.utc)))
                await db.commit()
        asyncio.run(seed())

        async def database():
            async with self.sessions() as db:
                yield db
        self._old = (settings.agency_workspace_id, routes.session_factory, routes.provider_factory,
                     routes.site_transport, routes.site_delay)
        # Авто-сбор контактов после поиска: в тестах сайты «недоступны», в сеть не ходим.
        routes.site_transport = httpx.MockTransport(lambda request: httpx.Response(404))
        routes.site_delay = 0
        settings.agency_workspace_id = 7
        routes.session_factory = self.sessions
        routes.provider_factory = lambda: StaticProvider(DATA)
        app = FastAPI(); app.include_router(api_router, prefix="/api")
        app.dependency_overrides[get_db] = database
        self.client = TestClient(app)
        self.client.cookies.set(COOKIE, "a" * 43)

    def tearDown(self):
        (settings.agency_workspace_id, routes.session_factory, routes.provider_factory,
         routes.site_transport, routes.site_delay) = self._old
        asyncio.run(self.engine.dispose())
        os.unlink(self.path)

    def post(self, url, json=None):
        return self.client.post(url, json=json, headers=ORIGIN)

    def run_search(self):
        r = self.post("/api/admin/leadgen/runs", {"keywords": KEYS, "region_code": 54, "niche": "Кухни на заказ",
                                                  "city": "Екатеринбург"})
        self.assertEqual(r.status_code, 202, r.text)
        return r.json()["id"]

    def test_requires_admin_and_origin(self):
        anon = TestClient(self.client.app)
        self.assertEqual(anon.get("/api/admin/leadgen/runs").status_code, 401)
        self.assertEqual(self.client.post("/api/admin/leadgen/runs", json={"keywords": KEYS}).status_code, 403)

    def test_workspace_not_configured(self):
        settings.agency_workspace_id = 0
        self.assertEqual(self.client.get("/api/admin/leadgen/runs").status_code, 503)

    def test_provider_not_configured(self):
        def broken():
            raise ProviderNotConfigured("XMLSTOCK_USER не задан")
        routes.provider_factory = broken
        r = self.post("/api/admin/leadgen/runs", {"keywords": KEYS})
        self.assertEqual(r.status_code, 400)
        self.assertIn("XMLSTOCK_USER", r.json()["detail"])

    def test_estimate(self):
        r = self.post("/api/admin/leadgen/runs/estimate", {"keywords": KEYS})
        self.assertEqual(r.json(), {"requests": 3, "cost_rub": 0.04})

    def test_run_then_results_and_company_card(self):
        run_id = self.run_search()
        run = self.client.get(f"/api/admin/leadgen/runs/{run_id}").json()
        self.assertEqual(run["status"], "done")
        self.assertNotIn("company_ids", run["stats"])
        self.assertEqual([c["domain"] for c in run["companies"]], ["kuhni-sever.ru", "dub-mebel.ru"])
        top = run["companies"][0]
        self.assertTrue(top["is_new"])
        self.assertEqual((top["score"], top["ad"]["placement"]), (7, "premium"))

        card = self.client.get(f"/api/admin/leadgen/companies/{top['id']}").json()
        self.assertEqual(card["ads"][0]["title"], "Кухни от 14 дней")
        self.assertIn("ad_direct", {s["kind"] for s in card["signals"]})
        self.assertEqual(self.client.get("/api/admin/leadgen/companies/9999").status_code, 404)
        self.assertEqual(len(self.client.get("/api/admin/leadgen/runs").json()["items"]), 1)
        by_run = self.client.get("/api/admin/leadgen/companies", params={"run_id": run_id}).json()
        self.assertEqual(by_run["total"], 2)

        # Выгрузка в Excel: те же фильтры, строка на компанию, контакты и объявление.
        import io
        from openpyxl import load_workbook
        r = self.client.get("/api/admin/leadgen/companies/export", params={"run_id": run_id})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("attachment", r.headers["content-disposition"])
        sheet = load_workbook(io.BytesIO(r.content)).active
        header = [c.value for c in sheet[1]]
        rows = [dict(zip(header, [c.value for c in row])) for row in sheet.iter_rows(min_row=2)]
        self.assertEqual([x["Сайт"] for x in rows], ["kuhni-sever.ru", "dub-mebel.ru"])
        self.assertEqual(rows[0]["Что рекламирует"], "Кухни от 14 дней")
        self.assertEqual(rows[0]["Спецразмещение"], "да")
        only = self.client.get("/api/admin/leadgen/companies/export", params={"ids": [top["id"]]})
        self.assertEqual(load_workbook(io.BytesIO(only.content)).active.max_row, 2)

    def test_companies_filters_and_sort(self):
        self.run_search()
        base = "/api/admin/leadgen/companies"
        self.assertEqual(self.client.get(base).json()["total"], 2)
        hot = self.client.get(base, params={"min_score": 6}).json()
        self.assertEqual([c["domain"] for c in hot["items"]], ["kuhni-sever.ru"])
        self.assertEqual(self.client.get(base, params={"signal": "ad_premium"}).json()["total"], 1)
        self.assertEqual(self.client.get(base, params={"q": "dub"}).json()["total"], 1)
        self.assertEqual(self.client.get(base, params={"city": "екатеринбург"}).json()["total"], 2)
        self.assertEqual(self.client.get(base, params={"channel": "telegram"}).json()["total"], 0)
        self.assertEqual(self.client.get(base, params={"stage": "bogus"}).status_code, 422)

    def test_segments(self):
        self.run_search()
        r = self.post("/api/admin/leadgen/segments", {"name": "Горячие кухни", "filters": {"min_score": 6}})
        self.assertEqual(r.status_code, 201, r.text)
        seg = r.json()
        self.assertEqual(seg["count"], 1)
        listed = self.client.get("/api/admin/leadgen/companies", params={"segment_id": seg["id"]}).json()
        self.assertEqual(listed["total"], 1)
        self.assertEqual(self.post("/api/admin/leadgen/segments",
                                   {"name": "x", "filters": {"price": 1}}).status_code, 422)
        upd = self.client.put(f"/api/admin/leadgen/segments/{seg['id']}", headers=ORIGIN,
                              json={"name": "Все", "filters": {}}).json()
        self.assertEqual(upd["count"], 2)
        self.assertEqual(self.client.delete(f"/api/admin/leadgen/segments/{seg['id']}", headers=ORIGIN).status_code, 204)

    def test_dnc_blocks_and_zeroes_score(self):
        self.run_search()
        sever = self.client.get("/api/admin/leadgen/companies", params={"q": "sever"}).json()["items"][0]
        ok = self.post(f"/api/admin/leadgen/companies/{sever['id']}/can-contact", {"channel": "email"}).json()
        self.assertTrue(ok["ok"])
        r = self.post("/api/admin/leadgen/dnc", {"kind": "domain", "value": "https://www.kuhni-sever.ru/",
                                                 "reason": "refused"})
        self.assertEqual((r.status_code, r.json()["value"]), (201, "kuhni-sever.ru"))
        blocked = self.post(f"/api/admin/leadgen/companies/{sever['id']}/can-contact", {"channel": "email"}).json()
        self.assertEqual(blocked, {"ok": False, "reason": "dnc:domain"})
        r = self.post("/api/admin/leadgen/dnc", {"kind": "company", "value": str(sever["id"]), "reason": "refused"})
        card = self.client.get(f"/api/admin/leadgen/companies/{sever['id']}").json()
        self.assertEqual((card["stage"], card["score"]), ("dnc", 0))
        items = self.client.get("/api/admin/leadgen/dnc").json()["items"]
        self.assertEqual(len(items), 2)
        self.assertEqual(self.client.delete(f"/api/admin/leadgen/dnc/{items[0]['id']}", headers=ORIGIN).status_code, 204)

    def test_merge_api(self):
        self.run_search()
        items = self.client.get("/api/admin/leadgen/companies").json()["items"]
        r = self.post("/api/admin/leadgen/companies/merge", {"keep_id": items[0]["id"], "drop_id": items[1]["id"]})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(self.client.get("/api/admin/leadgen/companies").json()["total"], 1)
        same = self.post("/api/admin/leadgen/companies/merge", {"keep_id": items[0]["id"], "drop_id": items[0]["id"]})
        self.assertEqual(same.status_code, 409)

    def test_settings_roundtrip(self):
        r = self.client.put("/api/admin/leadgen/settings", headers=ORIGIN,
                            json={"weights": {"ad_direct": 4}, "platform_domains": ["https://www.avito.ru", "zoon.ru"],
                                  "recontact_days": 60})
        self.assertEqual(r.status_code, 200, r.text)
        s = self.client.get("/api/admin/leadgen/settings").json()
        self.assertEqual((s["weights"]["ad_direct"], s["weights"]["has_messenger"]), (4, 2))
        self.assertEqual((s["platform_domains"], s["recontact_days"]), (["avito.ru", "zoon.ru"], 60))
        bad = self.client.put("/api/admin/leadgen/settings", headers=ORIGIN, json={"weights": {"ad_direct": 99}})
        self.assertEqual(bad.status_code, 422)


    def test_import_csv_and_xlsx(self):
        import io
        from openpyxl import Workbook
        base = "/api/admin/leadgen/companies/import"
        csv_text = ("Название;Сайт;Телефон;Email;Город\n"
                    "Ромакс;https://romax63.ru/;8 905 303-66-00;info@romax63.ru;Самара\n"
                    "Авито;avito.ru;;;\n"
                    "Без сайта;;;;\n"
                    "Ромакс дубль;www.romax63.ru;;;\n")
        r = self.client.post(base, headers=ORIGIN, data={"text": csv_text, "niche": "Кухни", "enrich": "false"})
        self.assertEqual(r.status_code, 201, r.text)
        st = r.json()["stats"]
        self.assertEqual((st["companies"], st["created"], st["invalid"]), (1, 1, 2))  # дубль склеен, Авито и пустая — мимо
        self.assertEqual(r.json()["enrich_queued"], 0)
        card = self.client.get("/api/admin/leadgen/companies", params={"run_id": r.json()["id"]}).json()
        self.assertEqual([c["domain"] for c in card["items"]], ["romax63.ru"])
        self.assertEqual(card["items"][0]["niche"], "Кухни")

        wb = Workbook(); ws = wb.active
        ws.append(["Сайт компании", "Телефон 1", "ИНН"])
        ws.append(["agata63.ru", "+7 846 111-22-33", None])
        ws.append(["romax63.ru", None, None])
        buf = io.BytesIO(); wb.save(buf)
        r = self.client.post(base, headers=ORIGIN, data={"enrich": "false"},
                             files={"file": ("baza.xlsx", buf.getvalue(), "application/octet-stream")})
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual((r.json()["stats"]["created"], r.json()["stats"]["already_in_base"]), (1, 1))
        self.assertEqual(r.json()["params"]["filename"], "baza.xlsx")
        bad = self.client.post(base, headers=ORIGIN, data={"text": "просто текст без сайтов"})
        self.assertEqual(bad.status_code, 422)
        old = self.client.post(base, headers=ORIGIN, files={"file": ("old.xls", b"x", "application/vnd.ms-excel")})
        self.assertEqual(old.status_code, 422)

    def test_hide_restore_delete_and_block(self):
        base = "/api/admin/leadgen"
        run_id = self.run_search()
        items = self.client.get(f"{base}/companies").json()["items"]
        ids = {c["domain"]: c["id"] for c in items}
        sever, dub = ids["kuhni-sever.ru"], ids["dub-mebel.ru"]

        r = self.post(f"{base}/companies/hide", {"company_ids": [sever]})
        self.assertEqual(r.json()["done"], 1)
        self.assertEqual([c["domain"] for c in self.client.get(f"{base}/companies").json()["items"]], ["dub-mebel.ru"])
        hidden = self.client.get(f"{base}/companies", params={"stage": "hidden"}).json()["items"]
        self.assertEqual([c["domain"] for c in hidden], ["kuhni-sever.ru"])
        # повторный поиск не возвращает скрытую компанию в список
        self.run_search()
        self.assertNotIn("kuhni-sever.ru", [c["domain"] for c in self.client.get(f"{base}/companies").json()["items"]])
        self.post(f"{base}/companies/hide", {"company_ids": [sever], "hidden": False})
        card = self.client.get(f"{base}/companies/{sever}").json()
        self.assertNotEqual(card["stage"], "hidden")

        r = self.post(f"{base}/companies/delete", {"company_ids": [dub], "block_domains": True})
        self.assertEqual((r.json()["deleted"], r.json()["blocked_domains"]), (1, 1))
        self.assertEqual(self.client.get(f"{base}/companies/{dub}").status_code, 404)
        self.assertIn("dub-mebel.ru", self.client.get(f"{base}/settings").json()["platform_domains"])
        self.run_search()  # сайт в исключениях — больше не добавляется
        self.assertNotIn("dub-mebel.ru", [c["domain"] for c in self.client.get(f"{base}/companies").json()["items"]])

        # компанию с идущей цепочкой не удаляем и не скрываем
        from app.domains.leadgen.models import LgEnrollment, LgSequence

        async def busy():
            async with self.sessions() as db:
                seq = LgSequence(workspace_id=7, name="x", steps=[], window={})
                db.add(seq); await db.flush()
                db.add(LgEnrollment(workspace_id=7, sequence_id=seq.id, company_id=sever, status="active",
                                    current_step=0))
                await db.commit()
        asyncio.run(busy())
        r = self.post(f"{base}/companies/delete", {"company_ids": [sever]})
        self.assertEqual((r.json()["deleted"], r.json()["skipped"][0]["reason_text"][:12]), (0, "идёт цепочка"))
        self.assertEqual(self.post(f"{base}/companies/hide", {"company_ids": [sever]}).json()["done"], 0)


class JobQueueTests(unittest.TestCase):
    def test_background_jobs_run_one_at_a_time(self):
        """Повторные нажатия «Собрать контакты» не запускают параллельные обходы сайтов."""
        state = {"now": 0, "max": 0, "order": []}
        original = routes._enrich_now

        async def fake(ids, *, only_stale):
            state["now"] += 1
            state["max"] = max(state["max"], state["now"])
            await asyncio.sleep(0.01)
            state["order"].append(ids[0])
            state["now"] -= 1

        async def go():
            await asyncio.gather(*(routes._enrich([i], only_stale=True) for i in (1, 2, 3)),
                                 routes._assess_now([], False))  # другая очередь — не ждёт
        routes._enrich_now = fake
        try:
            asyncio.run(go())
        finally:
            routes._enrich_now = original
        self.assertEqual((state["max"], sorted(state["order"])), (1, [1, 2, 3]))

if __name__ == "__main__":
    unittest.main()
